#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
NÜMKO · Actualizar catálogo
===========================
Lee  catalogo.xlsx  (hoja PRODUCTOS)  y genera  productos.js  que usa la página.

    EDITAR EXCEL  ->  ejecutar este script (o ACTUALIZAR_CATALOGO.bat)  ->  página actualizada

Reglas:
  * Un perfume se identifica por  MARCA + NOMBRE  (nunca por un número).
  * Si hay varias filas con la misma MARCA + NOMBRE + CATEGORÍA (por ejemplo, el mismo perfume en
    50 ML y 100 ML, en EDP y EXTRAIT, o en versión TESTER) se juntan en UNA sola tarjeta y las
    versiones se eligen en el detalle. Si dos filas son idénticas, se publica la primera y se avisa.
  * La foto se busca sola en la carpeta fotos/ comparando MARCA + NOMBRE con el nombre del
    archivo. Si quieres forzar una foto, escribe el nombre del archivo en la columna FOTO.
  * productos.js es un archivo GENERADO. No lo edites a mano: se sobrescribe cada vez.

Requisito único:  pip install openpyxl
"""
import sys, os, re, json, unicodedata, difflib, datetime
from collections import defaultdict

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Ejecuta:  pip install openpyxl")

RAIZ = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
EXCEL = os.path.join(RAIZ, "catalogo.xlsx")
HOJA = "PRODUCTOS"
CARPETA_FOTOS = os.path.join(RAIZ, "fotos")
SALIDA_JS = os.path.join(RAIZ, "productos.js")
SALIDA_REPORTE = os.path.join(RAIZ, "reporte_actualizacion.txt")
EXT_FOTO = [".jpg", ".jpeg", ".webp", ".png"]          # prioridad si hay dos archivos iguales
COLECCIONES = {"arabe", "disenador", "nicho", "revista", "aerosoles", "accesorios"}
CATEGORIAS = {"HOMBRE", "MUJER", "UNISEX"}

# ----------------------------------------------------------------------------- normalización
def sin_acentos(s):
    s = unicodedata.normalize("NFD", str(s))
    return "".join(c for c in s if unicodedata.category(c) != "Mn")

def norm(s):
    """MAYÚSCULAS, sin acentos ni símbolos: 'Dolce & Gabbana' -> 'DOLCE GABBANA'."""
    s = sin_acentos(s).upper().replace("Ü", "U")
    s = re.sub(r"[’'`´]", "", s)
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return s.strip()

def slug(s):
    s = sin_acentos(s).upper().replace("&", " AND ")
    s = re.sub(r"[’'`´]", "", s)
    return re.sub(r"[^A-Z0-9]+", "-", s).strip("-")

def clave_base(marca, nombre):
    return f"{slug(marca)}__{slug(nombre)}"

def limpiar(s):
    return re.sub(r"\s+", " ", str(s).strip()) if s is not None else ""

def a_numero(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    t = re.sub(r"[^\d,.\-]", "", str(v)).replace(".", "").replace(",", ".")
    try:
        return int(round(float(t)))
    except ValueError:
        return None

# ----------------------------------------------------------------------------- fotos
ALIAS = {"FRAGANCE": "FRAGRANCE", "LATAFFA": "LATTAFA"}          # errores de tipeo frecuentes
RUIDO = {"PERFUME", "ML", "SPRAY", "ESTUCHE", "PARA", "CLON", "DESODORANTE", "PERFUMES"}
OPCIONALES = {"AND", "Y", "DE", "EL", "LA", "BY", "THE", "S"}      # no se exigen al comparar
TIPOS = {"EDP", "EDT", "PARFUM", "EXTRAIT", "TESTER", "SET", "SPLASH", "MIST", "LOCION", "COLONIA"}
GENERO = {
    "HOMBRE": {"HOMBRE", "VARON", "MEN", "MAN", "MALE", "HIM"},
    "MUJER": {"MUJER", "DAMA", "WOMAN", "WOMEN", "FEMALE", "HER", "FEMME"},
    "UNISEX": {"UNISEX"},
}
TODO_GENERO = set().union(*GENERO.values())

def _decodifica(s):
    """Algunos programas guardan 'É' como '#U00c9' en el nombre del archivo: lo devolvemos a la letra real."""
    return re.sub(r"#U([0-9A-Fa-f]{4})", lambda m: chr(int(m.group(1), 16)), str(s))

def tokens(s):
    t = norm(_decodifica(s))
    t = re.sub(r"\b(\d+) (AM|PM)\b", r"\1\2", t)               # '9 PM' = '9PM'
    return [ALIAS.get(x, x) for x in t.split()]

def volumenes(s):
    n = norm(_decodifica(s))
    return {int(x) for x in re.findall(r"(\d+)\s*ML", n)} or {int(x) for x in re.findall(r"\b(\d{2,3})\b", n)}

def _parecidos(t, u):
    if t == u: return True
    if t.isdigit() or u.isdigit() or len(t) < 4 or len(u) < 4: return False
    return difflib.SequenceMatcher(None, t, u).ratio() >= 0.8

class Fotos:
    def __init__(self, carpeta):
        self.carpeta = carpeta
        self.archivos = {}          # nombre real -> info
        if os.path.isdir(carpeta):
            por_stem = {}
            for f in sorted(os.listdir(carpeta)):
                st, ext = os.path.splitext(f)
                if ext.lower() not in EXT_FOTO:
                    continue
                k = norm(_decodifica(st))
                prev = por_stem.get(k)
                if prev is None or EXT_FOTO.index(ext.lower()) < EXT_FOTO.index(os.path.splitext(prev)[1].lower()):
                    por_stem[k] = f
            for f in por_stem.values():
                st = os.path.splitext(f)[0]
                tk = tokens(st)
                self.archivos[f] = {"tok": tk, "set": set(tk), "vol": volumenes(st)}
        self._por_norm = {norm(_decodifica(os.path.splitext(f)[0])): f for f in self.archivos}

    def existe(self, nombre):
        if not nombre: return None
        st, ext = os.path.splitext(nombre)
        if ext.lower() not in EXT_FOTO: st = nombre
        return self._por_norm.get(norm(_decodifica(st)))

    # --- piezas de puntuación
    @staticmethod
    def _propios(K):
        return K | TIPOS | TODO_GENERO | RUIDO | OPCIONALES

    def _extras(self, K, info):
        return [t for t in info["tok"] if t not in self._propios(K) and not t.isdigit() and not re.fullmatch(r"\d+ML", t)]

    def puntaje(self, p, K, f, extras):
        info = self.archivos[f]
        pts = -3 * extras
        vp = volumenes(p["volumen"])
        if vp and info["vol"]:
            pts += 4 if vp & info["vol"] else -4
            if len(info["vol"] - vp) > 0 and vp & info["vol"]: pts -= 1       # la foto muestra más volúmenes (pack)
        tp = (set(tokens(p["tipo"])) & TIPOS) - K
        ft = (info["set"] & TIPOS) - K
        if tp and (tp & ft): pts += 2
        if "TESTER" in ft and "TESTER" not in tp: pts -= 3
        if "TESTER" in tp and "TESTER" not in ft: pts -= 1
        if "EXTRAIT" in ft and "EXTRAIT" not in tp: pts -= 2
        if "ESTUCHE" in info["set"] and "SET" not in tp and p["presentacion"].lower().startswith("perfume solo"): pts -= 2
        gf = info["set"] & TODO_GENERO
        if gf and p["categoria"] in ("HOMBRE", "MUJER"):
            pts += 1 if gf & GENERO[p["categoria"]] else -2
        return pts

    def _marca_al_inicio(self, p, info):
        mt = [t for t in tokens(p["marca"]) if t not in OPCIONALES]
        pos = [info["tok"].index(t) for t in mt if t in info["tok"]]
        return bool(pos) and min(pos) <= 2

    # --- fase 1: todas las palabras de MARCA + NOMBRE aparecen en el nombre del archivo
    def _fase1(self, productos, libres):
        info_p, min_extras = {}, defaultdict(lambda: 99)
        for i, p in enumerate(productos):
            if i not in libres: continue
            K = {t for t in tokens(p["marca"]) + tokens(p["nombre"]) if t not in OPCIONALES}
            cands = []
            for f, info in self.archivos.items():
                # en las fotos de 'clones' lo que va después de la palabra CLON es otro perfume: no cuenta
                tope = info["tok"].index("CLON") if "CLON" in info["tok"] else len(info["tok"])
                if K <= set(info["tok"][:tope]) and self._marca_al_inicio(p, info):
                    ex = len(self._extras(K, info))
                    cands.append((f, ex)); min_extras[f] = min(min_extras[f], ex)
            info_p[i] = (K, cands)
        res, amb = {}, {}
        for i, (K, cands) in info_p.items():
            cands = [(f, ex) for f, ex in cands if ex == min_extras[f]]      # la foto es de quien mejor la describe
            if not cands: continue
            pts = sorted(((self.puntaje(productos[i], K, f, ex), len(f), f) for f, ex in cands), key=lambda x: (-x[0], x[1], x[2]))
            mejores = [x for x in pts if x[0] == pts[0][0]]
            res[i] = (pts[0][2], "nombre" if len(mejores) == 1 else "varias")
            if len(mejores) > 1: amb[i] = [x[2] for x in mejores]
        return res, amb

    # --- fase 2: nombres con una letra mal escrita, o fotos sin la marca
    def _fase2(self, productos, libres, usadas):
        res = {}
        for i in libres:
            p = productos[i]
            nom = [t for t in tokens(p["nombre"]) if t not in OPCIONALES]
            mar = [t for t in tokens(p["marca"]) if t not in OPCIONALES]
            if not nom: continue
            vp = volumenes(p["volumen"])
            cand = []
            for f, info in self.archivos.items():
                if f in usadas: continue
                if vp and info["vol"] and not (vp & info["vol"]): continue
                ft = [t for t in info["tok"] if t not in OPCIONALES]
                usados = set()
                def hay(t):
                    for u in ft:
                        if _parecidos(t, u): usados.add(u); return True
                    return False
                if not all(hay(t) for t in nom): continue
                bm = sum(1 for t in mar if hay(t))
                if not (bm >= max(1, (len(mar) + 1) // 2) or (bm == 0 and len(nom) >= 2)): continue
                extras = [u for u in ft if u not in usados and u not in TIPOS | TODO_GENERO | RUIDO and not u.isdigit()
                          and not re.fullmatch(r"\d+ML", u)]
                if len(extras) > 2: continue
                K = set(nom + mar)
                cand.append((self.puntaje(p, K, f, len(extras)) - (3 if bm == 0 else 0), len(f), f))
            cand.sort(key=lambda x: (-x[0], x[1], x[2]))
            if cand and (len(cand) == 1 or cand[0][0] > cand[1][0]):
                res[i] = (cand[0][2], "parecido")
        return res

    def asignar(self, productos):
        """Devuelve ({indice: (archivo, metodo)}, {indice: [fotos igual de buenas]})."""
        todos = set(range(len(productos)))
        res, amb = self._fase1(productos, todos)
        res2 = self._fase2(productos, todos - set(res), {f for f, _ in res.values()})
        # una foto 'parecida' no se puede asignar a dos perfumes distintos a la vez
        cuenta = defaultdict(list)
        for i, (f, m) in res2.items(): cuenta[f].append(i)
        for f, idx in cuenta.items():
            if len(idx) == 1: res[idx[0]] = res2[idx[0]]
        return res, amb

# ----------------------------------------------------------------------------- lectura del Excel
COLS = {
    "MARCA": "marca", "NOMBRE": "nombre", "CATEGORIA": "categoria", "COLECCION": "coleccion",
    "TIPO": "tipo", "VOLUMEN": "volumen", "PRESENTACION": "presentacion", "PRECIO": "precio",
    "RETAIL": "retail", "PRECIO TESTER": "precio_tester", "RETAIL TESTER": "retail_tester",
    "NOTAS": "notas", "FOTO": "foto", "MOSTRAR": "mostrar", "NOMBRE ORIGINAL": "nombre_original",
}

def leer_excel(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    if HOJA not in wb.sheetnames:
        sys.exit(f"No encuentro la hoja '{HOJA}' en {os.path.basename(ruta)}.")
    ws = wb[HOJA]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        sys.exit("La hoja PRODUCTOS está vacía.")
    cab = {}
    for j, h in enumerate(filas[0]):
        if h is None: continue
        k = norm(re.sub(r"\(.*?\)", "", str(h)))
        if k in COLS: cab[COLS[k]] = j
    faltan = [c for c in ("MARCA", "NOMBRE", "PRECIO", "RETAIL", "COLECCION") if COLS[c] not in cab]
    if faltan:
        sys.exit("Faltan columnas obligatorias en la fila 1 de PRODUCTOS: " + ", ".join(faltan))
    out = []
    for n, f in enumerate(filas[1:], start=2):
        g = lambda k: (f[cab[k]] if k in cab and cab[k] < len(f) else None)
        if all(v in (None, "") for v in f):
            continue
        out.append({
            "fila": n,
            "marca": limpiar(g("marca")).upper(), "nombre": limpiar(g("nombre")).upper(),
            "categoria": limpiar(g("categoria")).upper(),
            "coleccion": sin_acentos(limpiar(g("coleccion"))).lower(),
            "tipo": limpiar(g("tipo")).upper(), "volumen": limpiar(g("volumen")).upper(),
            "presentacion": limpiar(g("presentacion")),
            "precio": a_numero(g("precio")), "retail": a_numero(g("retail")),
            "precio_tester": a_numero(g("precio_tester")), "retail_tester": a_numero(g("retail_tester")),
            "notas": limpiar(g("notas")), "foto": limpiar(g("foto")),
            "mostrar": sin_acentos(limpiar(g("mostrar"))).upper() or "SI",
            "nombre_original": limpiar(g("nombre_original")),
        })
    return out

# ----------------------------------------------------------------------------- principal
def main():
    if not os.path.exists(EXCEL):
        sys.exit(f"No encuentro {EXCEL}")
    todas = leer_excel(EXCEL)
    errores, avisos = [], []

    ocultas = [p for p in todas if p["mostrar"] in ("NO", "N", "OCULTO", "0")]
    prods = [p for p in todas if p not in ocultas]

    # validaciones
    for p in prods:
        e = f"fila {p['fila']} ({p['marca']} {p['nombre']})"
        if not p["marca"] or not p["nombre"]: errores.append(f"{e}: falta MARCA o NOMBRE")
        if p["precio"] is None: errores.append(f"{e}: PRECIO vacío o no es número")
        if p["retail"] is None: errores.append(f"{e}: RETAIL vacío o no es número")
        if p["coleccion"] not in COLECCIONES: errores.append(f"{e}: COLECCIÓN '{p['coleccion']}' no válida (usa: {', '.join(sorted(COLECCIONES))})")
        if p["categoria"] not in CATEGORIAS: avisos.append(f"{e}: CATEGORÍA '{p['categoria']}' (se esperaba HOMBRE, MUJER o UNISEX)")
        if p["precio"] is not None and p["retail"] is not None and p["retail"] < p["precio"]:
            avisos.append(f"{e}: el RETAIL ({p['retail']}) es menor que el PRECIO ({p['precio']})")
        if p["precio_tester"] is not None and (p["retail_tester"] is None or p["retail_tester"] <= p["precio_tester"]):
            avisos.append(f"{e}: el tester no tiene descuento (precio tester {p['precio_tester']}, retail tester {p['retail_tester']})")
    # ---- agrupar: un perfume = MARCA + NOMBRE (+ categoría). Cada fila del Excel es una VERSIÓN
    #      (otro volumen, extrait, tester...) que se elige dentro del detalle.
    def vol_norm(v): return re.sub(r"\s+", "", v or "")
    def num_vol(v):
        m = re.search(r"(\d+)", v or "")
        return int(m.group(1)) if m else 0
    def es_tester(p): return "TESTER" in tokens(p["tipo"])
    def orden_version(v): return (bool(v.get("tester")), num_vol(v["volumen"]), v["tipo"])

    grupos, orden = {}, []
    for p in prods:
        k = (clave_base(p["marca"], p["nombre"]), p["categoria"])
        if k not in grupos: grupos[k] = []; orden.append(k)
        grupos[k].append(p)

    # filas idénticas (mismo perfume, tipo, volumen y presentación): se publica la primera
    repetidas, descartadas = [], set()
    for g in grupos.values():
        vistos = {}
        for p in g:
            firma = (es_tester(p), slug(p["tipo"]), vol_norm(p["volumen"]), (p["presentacion"] or "Perfume solo").lower())
            if firma in vistos:
                repetidas.append(f"{p['marca']} {p['nombre']} {p['tipo']} {p['volumen']}: se publicó la fila {vistos[firma]['fila']} "
                                 f"(precio {vistos[firma]['precio']}) y NO la fila {p['fila']} (precio {p['precio']}). "
                                 f"Borra la que sobra o escribe NO en MOSTRAR.")
                descartadas.add(id(p))
            else:
                vistos[firma] = p
    unicas = [p for p in prods if id(p) not in descartadas]

    # claves: MARCA + NOMBRE; si el mismo nombre existe en dos categorías, se agrega la categoría
    por_base = defaultdict(list)
    for k in orden: por_base[k[0]].append(k)
    clave_de, categorias_distintas = {}, []
    for base, ks in por_base.items():
        for k in ks:
            clave_de[k] = base if len(ks) == 1 else f"{base}__{slug(k[1]) or 'SIN-CATEGORIA'}"
        if len(ks) > 1:
            det = " / ".join(f"{k[1] or 'sin categoría'} (fila {grupos[k][0]['fila']})" for k in ks)
            categorias_distintas.append(f"{grupos[ks[0]][0]['marca']} {grupos[ks[0]][0]['nombre']}: {det}")

    if errores:
        escribir_reporte(errores, avisos, [], [], [], [], [], len(todas), 0, fallo=True)
        print("\nNO SE ACTUALIZÓ NADA. Corrige esto en el Excel y vuelve a ejecutar:\n  - " + "\n  - ".join(errores))
        sys.exit(1)

    # fotos (se busca la foto de cada versión)
    fotos = Fotos(CARPETA_FOTOS)
    auto, amb = fotos.asignar(unicas)
    foto_mala, por_parecido, elegidas = [], [], []
    usadas = defaultdict(set)
    for i, p in enumerate(unicas):
        archivo = ""
        if p["foto"]:
            real = fotos.existe(p["foto"])
            if real: archivo = real
            else: foto_mala.append(f"fila {p['fila']} ({p['marca']} {p['nombre']}): la FOTO '{p['foto']}' no existe en fotos/")
        if not archivo and i in auto:
            archivo, metodo = auto[i]
            if metodo == "parecido": por_parecido.append(f"fila {p['fila']} {p['marca']} {p['nombre']} {p['volumen']}  →  {archivo}")
            if metodo == "varias": elegidas.append(f"fila {p['fila']} {p['marca']} {p['nombre']} {p['tipo']} {p['volumen']}: elegí '{archivo}'; también servían: " + " | ".join(x for x in amb[i] if x != archivo))
        p["_foto"] = archivo
        if archivo: usadas[archivo].add(clave_de[(clave_base(p["marca"], p["nombre"]), p["categoria"])])

    # salida productos.js: un objeto por perfume, con su lista de 'variantes'
    db, sin_foto, compartidas_idx = [], [], defaultdict(list)
    for k in orden:
        filas = [p for p in grupos[k] if id(p) not in descartadas]
        p0 = filas[0]
        variantes = []
        for p in filas:
            v = {"tipo": p["tipo"], "volumen": p["volumen"], "presentacion": p["presentacion"] or "Perfume solo",
                 "precio": p["precio"], "precioRetail": p["retail"], "foto": p["_foto"]}
            if es_tester(p): v["tester"] = True
            variantes.append(v)
        for p in filas:        # testers declarados en las columnas PRECIO TESTER / RETAIL TESTER
            if p["precio_tester"] is not None and p["retail_tester"] is not None and not any(
                    v.get("tester") and vol_norm(v["volumen"]) == vol_norm(p["volumen"]) for v in variantes):
                variantes.append({"tipo": "TESTER", "volumen": p["volumen"], "presentacion": "Tester",
                                  "precio": p["precio_tester"], "precioRetail": p["retail_tester"], "foto": "", "tester": True})
        variantes.sort(key=orden_version)
        con_foto = [v for v in variantes if v["foto"] and not v.get("tester")] or [v for v in variantes if v["foto"]]
        foto_grupo = max(con_foto, key=lambda v: num_vol(v["volumen"]))["foto"] if con_foto else ""
        for v in variantes:
            if v["foto"] == foto_grupo: v["foto"] = ""        # misma foto que el perfume: no se repite
        colecciones = {p["coleccion"] for p in filas}
        if len(colecciones) > 1:
            avisos.append(f"{p0['marca']} {p0['nombre']}: sus versiones tienen distinta COLECCIÓN ({', '.join(sorted(colecciones))}); se usó '{p0['coleccion']}'")
        if not foto_grupo:
            sin_foto.append(f"fila {p0['fila']} {p0['marca']} {p0['nombre']} ({', '.join((v['tipo'] + ' ' + v['volumen']).strip() for v in variantes)})")
        db.append({
            "key": clave_de[k], "marca": p0["marca"], "nombre": p0["nombre"],
            "nombreOriginal": p0["nombre_original"] or f"{p0['marca']} {p0['nombre']}",
            "categoria": p0["categoria"], "coleccion": p0["coleccion"],
            "notas": [n.strip() for n in next((p["notas"] for p in filas if p["notas"]), "").split(",") if n.strip()],
            "foto": foto_grupo, "variantes": variantes,
        })
    ahora = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    n_versiones = sum(len(d["variantes"]) for d in db)
    with open(SALIDA_JS, "w", encoding="utf-8") as fh:
        fh.write("// ============================================================================\n"
                 "// productos.js — ARCHIVO GENERADO AUTOMÁTICAMENTE. NO LO EDITES A MANO.\n"
                 "// La fuente del catálogo es catalogo.xlsx. Para cambiar un perfume:\n"
                 "//   1) edita catalogo.xlsx   2) ejecuta ACTUALIZAR_CATALOGO   3) recarga la página\n"
                 f"// Generado: {ahora} · {len(db)} perfumes · {n_versiones} versiones\n"
                 "// Cada perfume se identifica por su 'key' = MARCA + NOMBRE (sin números de ID)\n"
                 "// y trae sus 'variantes' (volumen, extrait, tester...) para elegir en el detalle.\n"
                 "// ============================================================================\n")
        fh.write("const PRODUCTOS_DB = ")
        fh.write(json.dumps(db, ensure_ascii=False, indent=0, separators=(",", ":")).replace("},{", "},\n{"))
        fh.write(";\n")

    # posibles repetidos con el nombre escrito distinto (misma categoría): solo para revisar
    SOBRAN = {"MAN", "MEN", "WOMAN", "WOMEN", "HOMBRE", "MUJER", "HER", "HIM", "FOR", "POUR", "HOMME", "FEMME",
              "UNISEX", "TESTER", "EDP", "EDT", "PARFUM", "EXTRAIT", "ML"}
    MASC, FEM = {"MAN", "MEN", "HOMBRE", "HIM", "HOMME"}, {"WOMAN", "WOMEN", "MUJER", "HER", "FEMME"}
    def compacto(n): return "".join(t for t in norm(n).split() if t not in SOBRAN and not re.fullmatch(r"\d*ML", t))
    similares = defaultdict(set)
    for d in db:
        c = compacto(d["nombre"])
        if c: similares[(d["marca"], c, d["categoria"])].add(d["nombre"])
    def genero(n):
        t = set(norm(n).split())
        return (bool(t & MASC), bool(t & FEM))
    parecidos = []
    for (m, c, cat), ns in similares.items():
        if len(ns) < 2: continue
        gs = [genero(n) for n in ns]
        if any(g == (True, False) for g in gs) and any(g == (False, True) for g in gs): continue   # hombre vs mujer: son distintos
        parecidos.append(f"{m}: " + "  /  ".join(sorted(ns)))

    # fotos sin producto / compartidas entre perfumes distintos
    asignadas = {p["_foto"] for p in unicas if p["_foto"]}
    huerfanas = sorted(f for f in fotos.archivos if f not in asignadas)
    compartidas = [(f, sorted(ks)) for f, ks in usadas.items() if len(ks) > 1]
    escribir_reporte(errores, avisos, sin_foto + foto_mala, huerfanas, compartidas, por_parecido, ocultas,
                     len(todas), len(db), fotos=len(fotos.archivos), con_foto=len(db) - len(sin_foto), elegidas=elegidas,
                     versiones=n_versiones, repetidas=repetidas, categorias=categorias_distintas, parecidos=parecidos)
    print(f"\nListo: {len(db)} perfumes ({n_versiones} versiones) → productos.js  ({len(db) - len(sin_foto)} con foto, {len(sin_foto)} sin foto)")
    print(f"Detalle y avisos en: {os.path.basename(SALIDA_REPORTE)}")

def escribir_reporte(errores, avisos, sin_foto, huerfanas, compartidas, parecido, ocultas, total_filas, total_web,
                     fallo=False, fotos=0, con_foto=0, elegidas=(), versiones=0, repetidas=(), categorias=(), parecidos=()):
    L = [f"REPORTE DE ACTUALIZACIÓN · {datetime.datetime.now():%Y-%m-%d %H:%M}", "=" * 70]
    if fallo:
        L += ["NO SE ACTUALIZÓ productos.js (hay errores que corregir en el Excel):", ""] + [f" - {e}" for e in errores]
    else:
        L += [f"Filas en el Excel: {total_filas}", f"Perfumes publicados en la web: {total_web} (con {versiones} versiones en total)",
              f"Filas ocultas (MOSTRAR = NO): {len(ocultas)}", f"Fotos disponibles en fotos/: {fotos}",
              f"Perfumes con foto: {con_foto}", f"Perfumes SIN foto: {len(sin_foto)}", ""]
    def sec(t, items, tope=None):
        if not items: return
        L.extend(["", f"## {t} ({len(items)})"])
        for x in (items if tope is None else items[:tope]): L.append(f" - {x if isinstance(x, str) else x}")
    sec("AVISOS DE DATOS (revisar, no bloquean)", avisos)
    sec("Filas repetidas: misma versión escrita dos veces (solo se publicó una)", list(repetidas))
    sec("Mismo nombre con distinta CATEGORÍA: salen como perfumes separados (si es el mismo, unifica la categoría en el Excel)", list(categorias))
    sec("Nombres parecidos de la misma marca (si es el mismo perfume, escríbelos igual en el Excel para que se junten)", list(parecidos))
    sec("Perfumes SIN foto (escribe el nombre del archivo en la columna FOTO o agrega la foto a fotos/)", sin_foto)
    sec("Foto asignada por parecido de nombre (confirma que sea la correcta)", parecido)
    sec("Había varias fotos posibles: se eligió la más parecida (cambia la columna FOTO si prefieres otra)", list(elegidas))
    sec("Fotos que no coinciden con ningún perfume del catálogo", huerfanas)
    sec("Fotos compartidas por varios perfumes distintos", [f"{f}: " + " / ".join(ks) for f, ks in compartidas])
    sec("Filas ocultas (MOSTRAR = NO)", [f"fila {p['fila']}: {p['marca']} {p['nombre']} {p['tipo']} {p['volumen']} (precio {p['precio']})" for p in ocultas])
    with open(SALIDA_REPORTE, "w", encoding="utf-8") as fh: fh.write("\n".join(L) + "\n")

if __name__ == "__main__":
    main()
